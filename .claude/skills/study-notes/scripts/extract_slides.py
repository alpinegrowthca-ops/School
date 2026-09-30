#!/usr/bin/env python3
"""Extract lecture slides (PPTX or PDF) to Markdown, optionally rendering page images.

Usage:
  extract_slides.py lecture.pptx                 # Markdown to stdout (text, tables, speaker notes)
  extract_slides.py lecture.pdf -o L05_slides.md
  extract_slides.py lecture.pdf --images out/    # also render slide-001.png ... for viewing diagrams

Dependencies (install what you need):
  pip install python-pptx   # for .pptx
  pip install pymupdf       # for .pdf text and --images
"""
import argparse
import os
import sys


def _need(pkg, pip_name):
    sys.exit(f"Missing dependency '{pkg}'. Install with: pip install {pip_name}")


def extract_pptx(path):
    try:
        from pptx import Presentation
    except ImportError:
        _need("pptx", "python-pptx")

    prs = Presentation(path)
    out = [f"# Slides: {os.path.basename(path)}", ""]
    for i, slide in enumerate(prs.slides, start=1):
        title = None
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip() or None
        out.append(f"## Slide {i}" + (f": {title}" if title else ""))

        for shape in _iter_shapes(slide.shapes):
            if shape == slide.shapes.title:
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                out.extend(_table_md(shape.table))
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        out.append("  " * para.level + f"- {text}")
            elif shape.shape_type == 13:  # PICTURE
                alt = _alt_text(shape)
                out.append(f"- [image{': ' + alt if alt else ''}]")
            elif getattr(shape, "has_chart", False) and shape.has_chart:
                out.append("- [chart]")

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append("")
                out.append("**Speaker notes:**")
                out.extend(f"> {line}" if line.strip() else ">" for line in notes.splitlines())
        out.append("")
    return "\n".join(out)


def _iter_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == 6:  # GROUP
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _alt_text(shape):
    try:
        return shape._element.nvPicPr.cNvPr.get("descr", "").strip()
    except AttributeError:
        return ""


def _table_md(table):
    rows = [[cell.text.strip().replace("\n", " ").replace("|", "\\|") for cell in row.cells] for row in table.rows]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["", "| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    lines.append("")
    return lines


def _pymupdf():
    try:
        import pymupdf
    except ImportError:
        try:
            import fitz as pymupdf
        except ImportError:
            _need("pymupdf", "pymupdf")
    return pymupdf


def _open_pdf(path):
    return _pymupdf().open(path)


def extract_pdf(path):
    doc = _open_pdf(path)
    out = [f"# Slides: {os.path.basename(path)}", ""]
    sparse = 0
    for i, page in enumerate(doc, start=1):
        text = page.get_text("text").strip()
        out.append(f"## Slide {i}")
        if len(text) < 20:
            sparse += 1
            out.append("- [little or no extractable text; view the rendered image]")
        else:
            out.extend(line.rstrip() for line in text.splitlines() if line.strip())
        if page.get_images():
            out.append(f"- [{len(page.get_images())} embedded image(s)]")
        out.append("")
    if sparse > len(doc) / 3:
        out.append(f"> Note: {sparse}/{len(doc)} pages have little text (scanned or image-based). "
                   "Re-run with --images and read the PNGs.")
    return "\n".join(out)


def render_images(path, outdir, zoom):
    os.makedirs(outdir, exist_ok=True)
    if path.lower().endswith(".pptx"):
        sys.exit("Rendering .pptx needs a PDF first: soffice --headless --convert-to pdf <file>.pptx, "
                 "then re-run with the PDF.")
    pymupdf = _pymupdf()
    doc = pymupdf.open(path)
    mat = pymupdf.Matrix(zoom, zoom)
    for i, page in enumerate(doc, start=1):
        target = os.path.join(outdir, f"slide-{i:03d}.png")
        page.get_pixmap(matrix=mat).save(target)
    print(f"Rendered {len(doc)} page image(s) to {outdir}", file=sys.stderr)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", help=".pptx or .pdf")
    ap.add_argument("-o", "--out", help="write Markdown here instead of stdout")
    ap.add_argument("--images", metavar="DIR", help="render each PDF page to DIR/slide-NNN.png")
    ap.add_argument("--zoom", type=float, default=1.5, help="render scale for --images (default 1.5)")
    args = ap.parse_args()

    if not os.path.exists(args.file):
        sys.exit(f"File not found: {args.file}")
    ext = os.path.splitext(args.file)[1].lower()
    if ext == ".pptx":
        md = extract_pptx(args.file)
    elif ext == ".pdf":
        md = extract_pdf(args.file)
    else:
        sys.exit(f"Unsupported file type '{ext}'. Use .pptx or .pdf (convert .ppt/.key to one of these first).")

    if args.out:
        with open(args.out, "w", encoding="utf-8") as f:
            f.write(md)
        print(f"Wrote {args.out}", file=sys.stderr)
    else:
        print(md)

    if args.images:
        render_images(args.file, args.images, args.zoom)


if __name__ == "__main__":
    main()
